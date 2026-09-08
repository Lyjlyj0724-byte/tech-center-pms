# -*- coding: utf-8 -*-
"""生成工序模板导入 Excel：填写说明 + 工序模板(示例) + 项目清单(示例)"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
REQ_FONT = Font(color="C00000", size=10)  # 必填说明红色
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# ---------- Sheet 0: 填写说明 ----------
ws = wb.active
ws.title = "填写说明"
ws.column_dimensions["A"].width = 110
lines = [
    ("技术中心项目管理系统 · 工序模板导入表", True),
    ("", False),
    ("【填写步骤】", True),
    ("1. 在「工序模板」Sheet 中，按模板逐行填写工序。同一种工序路线（模板）的若干行，填相同的模板编号。", False),
    ("2. 如有在建项目需要一并导入，在「项目清单」Sheet 中填写，并为每个项目指定要套用的工序模板编号。", False),
    ("3. 填好后整表交给系统管理员导入。系统会逐行校验，有错会指出具体行号和原因，全部通过后才会写入。", False),
    ("", False),
    ("【关键规则】", True),
    ("· 黄色行为示例，正式导入前请删除示例行。", False),
    ("· 列名带 * 为必填。", False),
    ("· 工序序号：每个模板内从 1 开始连续编号（1、2、3…），决定默认加工顺序。", False),
    ("· 工序编码：同一模板内不可重复；建议用 OP10、OP20…跳号，便于以后在中间插入工序。", False),
    ("· 前置工序编码：不填表示第一道工序；多道工序并行前置时用英文逗号分隔，如 OP10,OP20；只能引用本模板内序号更小的工序。", False),
    ("· 标准工时(h)：用于按工时加权计算项目整体进度，填正数；不填则该工序按等权参与计算。", False),
    ("· 日期格式：YYYY-MM-DD，如 2026-10-01；计划完工日期不得早于计划开始日期。", False),
    ("· 项目清单中的「工序模板编号」必须已在「工序模板」Sheet 中定义。", False),
]
r = 1
for text, is_head in lines:
    c = ws.cell(row=r, column=1, value=text)
    c.alignment = WRAP
    if is_head:
        c.font = Font(bold=True, size=14 if r == 1 else 12)
    r += 1

# ---------- Sheet 1: 工序模板 ----------
ws1 = wb.create_sheet("工序模板")
headers1 = [
    ("模板编号 *", "同一模板填相同编号"),
    ("模板名称 *", "同一编号下名称一致"),
    ("工序序号 *", "从1连续编号"),
    ("工序编码 *", "模板内唯一，建议OP10/OP20跳号"),
    ("工序名称 *", "不超过50字"),
    ("标准工时(h)", "正数，可空"),
    ("前置工序编码", "可空，多个用英文逗号分隔"),
    ("建议设备类型", "可空"),
    ("备注", "可空"),
]
widths1 = [16, 22, 10, 12, 18, 12, 16, 14, 20]
for i, ((name, tip), w) in enumerate(zip(headers1, widths1), 1):
    c = ws1.cell(row=1, column=i, value=name)
    c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, THIN
    t = ws1.cell(row=2, column=i, value=tip)
    t.font, t.alignment, t.border = REQ_FONT, CENTER, THIN
    ws1.column_dimensions[get_column_letter(i)].width = w

examples1 = [
    ["TMPL-MECH-01", "机械加工标准路线", 1, "OP10", "下料", 2, "", "锯床", ""],
    ["TMPL-MECH-01", "机械加工标准路线", 2, "OP20", "粗加工", 8, "OP10", "数控车床", ""],
    ["TMPL-MECH-01", "机械加工标准路线", 3, "OP30", "热处理", 4, "OP20", "热处理炉", "外协可标此列"],
    ["TMPL-MECH-01", "机械加工标准路线", 4, "OP40", "精加工", 10, "OP30", "加工中心", ""],
    ["TMPL-MECH-01", "机械加工标准路线", 5, "OP50", "检验", 2, "OP40", "三坐标", ""],
]
for row in examples1:
    ws1.append(row)
for row in ws1.iter_rows(min_row=3, max_row=2 + len(examples1), max_col=len(headers1)):
    for c in row:
        c.fill, c.border, c.alignment = EXAMPLE_FILL, THIN, WRAP
ws1.freeze_panes = "A3"

# ---------- Sheet 2: 项目清单 ----------
ws2 = wb.create_sheet("项目清单")
headers2 = [
    ("项目编号 *", "唯一"),
    ("项目名称 *", ""),
    ("负责人 *", ""),
    ("优先级 *", "高 / 中 / 低"),
    ("计划开始日期 *", "YYYY-MM-DD"),
    ("计划完工日期 *", "不早于开始日期"),
    ("工序模板编号 *", "须已在「工序模板」Sheet 定义"),
]
widths2 = [16, 26, 12, 10, 16, 16, 18]
for i, ((name, tip), w) in enumerate(zip(headers2, widths2), 1):
    c = ws2.cell(row=1, column=i, value=name)
    c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, THIN
    t = ws2.cell(row=2, column=i, value=tip)
    t.font, t.alignment, t.border = REQ_FONT, CENTER, THIN
    ws2.column_dimensions[get_column_letter(i)].width = w

examples2 = [
    ["PRJ-2026-001", "XX 型号传动箱试制", "张三", "高", "2026-10-01", "2026-12-31", "TMPL-MECH-01"],
]
for row in examples2:
    ws2.append(row)
for row in ws2.iter_rows(min_row=3, max_row=2 + len(examples2), max_col=len(headers2)):
    for c in row:
        c.fill, c.border, c.alignment = EXAMPLE_FILL, THIN, WRAP
ws2.freeze_panes = "A3"

import os
out = "project-management-system/templates/工序模板导入模板.xlsx"
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print("saved:", out)

# 回读验证
from openpyxl import load_workbook
wb2 = load_workbook(out)
print("sheets:", wb2.sheetnames)
for name in wb2.sheetnames:
    s = wb2[name]
    print(f"  {name}: {s.max_row} rows x {s.max_column} cols")
print("工序模板 row3:", [c.value for c in wb2["工序模板"][3]])
print("项目清单 row3:", [c.value for c in wb2["项目清单"][3]])
